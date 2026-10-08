import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

import numpy as np
import soundfile as sf
from docx import Document

from app_logic.epub_reader import EpubChapter, EpubNavigation
from app_logic.pdf_reader import PdfChapter, PdfPage
from tests.pdf_fixture import write_blank_pdf
from app_logic import workspace as document_workspace
from webui.errors import VoxBenchError


def write_epub(path: Path) -> None:
    container = """<?xml version="1.0"?>
<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
  <rootfiles>
    <rootfile full-path="EPUB/package.opf"
      media-type="application/oebps-package+xml"/>
  </rootfiles>
</container>"""
    package = """<?xml version="1.0"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0">
  <manifest>
    <item id="nav" href="nav.xhtml" media-type="application/xhtml+xml"
      properties="nav"/>
    <item id="opening" href="chapters/opening.xhtml"
      media-type="application/xhtml+xml"/>
    <item id="second" href="chapters/second%20chapter.xhtml"
      media-type="application/xhtml+xml"/>
  </manifest>
  <spine>
    <itemref idref="nav"/>
    <itemref idref="opening"/>
    <itemref idref="second"/>
  </spine>
</package>"""
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr(
            "mimetype",
            "application/epub+zip",
            compress_type=zipfile.ZIP_STORED,
        )
        archive.writestr("META-INF/container.xml", container)
        archive.writestr("EPUB/package.opf", package)
        archive.writestr(
            "EPUB/nav.xhtml",
            "<html><body><nav>Contents</nav></body></html>",
        )
        archive.writestr(
            "EPUB/chapters/opening.xhtml",
            "<html><body><h1>Opening</h1>"
            "<p>Hello document world.</p><script>unsafe()</script></body></html>",
        )
        archive.writestr(
            "EPUB/chapters/second chapter.xhtml",
            "<html><body><h2>Next chapter</h2><p>Second section.</p></body></html>",
        )


class DocumentWorkspaceTests(unittest.TestCase):
    def test_pdf_import_keeps_one_editable_section_per_page(self):
        pages = [
            PdfPage(
                1,
                "Preface text.\nChapter One\nFirst chapter.\nChapter Two\nSecond chapter.",
            )
        ]
        chapters = [
            PdfChapter("Chapter One", 0, 1),
            PdfChapter("Chapter Two", 0, 1),
        ]
        with patch.object(document_workspace, "read_pdf_document", return_value=(pages, chapters)):
            sections = document_workspace._extract_pdf(Path("book.pdf"))

        self.assertEqual([section["title"] for section in sections], ["Page 1"])
        self.assertIn("Chapter One", sections[0]["text"])
        self.assertIn("Chapter Two", sections[0]["text"])

    def test_pdf_bookmarks_group_output_at_selected_depth(self):
        with tempfile.TemporaryDirectory() as directory:
            documents = Path(directory) / "documents"
            with patch.object(document_workspace, "DOCUMENTS_DIR", documents):
                document_id = "bookmarks"
                (documents / document_id).mkdir(parents=True)
                section = document_workspace._new_section(
                    "Page 1",
                    "Preface\n1. Introduction\nText\n1.1 Overview\nMore\n1.1.1 Detail\nDeep text",
                    source_page=1,
                )
                document_workspace.save_section(document_id, section)
                document_workspace.save_manifest(
                    {
                        "id": document_id,
                        "source_type": ".pdf",
                        "sections": [section["id"]],
                        "pdf_bookmarks": [
                            {"title": "Introduction", "level": 0, "page": 1},
                            {"title": "Overview", "level": 1, "page": 1},
                            {"title": "Detail", "level": 2, "page": 1},
                        ],
                    }
                )

                grouped = document_workspace.document_generation_groups(
                    document_id,
                    [section["id"]],
                    use_pdf_bookmarks=True,
                    bookmark_depth=1,
                )
                pages = document_workspace.document_generation_groups(
                    document_id,
                    [section["id"]],
                    use_pdf_bookmarks=False,
                    bookmark_depth=1,
                )

        self.assertEqual([group["title"] for group in grouped], [
            "Page 1", "Introduction", "Overview",
        ])
        self.assertEqual(
            [group["is_bookmark"] for group in grouped], [False, True, True]
        )
        self.assertIn(
            "1.1.1 Detail",
            "\n".join(part["text"] for part in grouped[-1]["parts"]),
        )
        self.assertEqual([group["title"] for group in pages], ["Page 1"])

    def test_epub_navigation_groups_output_at_selected_depth(self):
        with tempfile.TemporaryDirectory() as directory:
            documents = Path(directory) / "documents"
            with patch.object(document_workspace, "DOCUMENTS_DIR", documents):
                document_id = "epub_navigation"
                (documents / document_id).mkdir(parents=True)
                sections = [
                    document_workspace._new_section(
                        "Opening", "Opening\nFirst chapter.",
                        chapter_title="Opening", chapter_level=0,
                    ),
                    document_workspace._new_section(
                        "Part one", "Part one\nSecond chapter.",
                        chapter_title="Part one", chapter_level=1,
                    ),
                    document_workspace._new_section(
                        "Detail", "Detail\nNested text.",
                        chapter_title="Detail", chapter_level=2,
                    ),
                ]
                for section in sections:
                    document_workspace.save_section(document_id, section)
                document_workspace.save_manifest(
                    {
                        "id": document_id,
                        "source_type": ".epub",
                        "sections": [section["id"] for section in sections],
                    }
                )

                grouped = document_workspace.document_generation_groups(
                    document_id,
                    [section["id"] for section in sections],
                    use_pdf_bookmarks=True,
                    bookmark_depth=1,
                )

        self.assertEqual([group["title"] for group in grouped], ["Opening", "Part one"])
        self.assertTrue(all(group["is_bookmark"] for group in grouped))
        self.assertEqual(grouped[0]["parts"][0]["text"], "First chapter.")
        self.assertIn("Detail", grouped[1]["parts"][1]["text"])

    def test_chapterless_epub_and_pdf_keep_page_groups(self):
        with tempfile.TemporaryDirectory() as directory:
            documents = Path(directory) / "documents"
            with patch.object(document_workspace, "DOCUMENTS_DIR", documents):
                for source_type in (".epub", ".pdf"):
                    document_id = source_type[1:]
                    (documents / document_id).mkdir(parents=True)
                    sections = [
                        document_workspace._new_section("One", "First."),
                        document_workspace._new_section("Two", "Second."),
                    ]
                    for section in sections:
                        document_workspace.save_section(document_id, section)
                    document_workspace.save_manifest(
                        {
                            "id": document_id,
                            "source_type": source_type,
                            "sections": [section["id"] for section in sections],
                        }
                    )
                    grouped = document_workspace.document_generation_groups(
                        document_id,
                        [section["id"] for section in sections],
                        use_pdf_bookmarks=True,
                        bookmark_depth=1,
                    )
                    self.assertEqual([group["title"] for group in grouped], ["One", "Two"])
                    self.assertTrue(all(not group["is_bookmark"] for group in grouped))

    def test_detects_only_strong_pdf_table_of_contents_pages(self):
        with tempfile.TemporaryDirectory() as directory:
            documents = Path(directory) / "documents"
            with patch.object(document_workspace, "DOCUMENTS_DIR", documents):
                document_id = "toc"
                (documents / document_id).mkdir(parents=True)
                toc = document_workspace._new_section(
                    "Page 1",
                    "Introduction . . . 2\nUser Guide . . . 3\nDeveloper Guide . . . 4",
                    source_page=1,
                )
                chapter = document_workspace._new_section(
                    "Page 2", "Introduction\nNormal chapter text.", source_page=2
                )
                toc["text"] = toc["text"].replace("\n", " ")
                for section in (toc, chapter):
                    document_workspace.save_section(document_id, section)
                document_workspace.save_manifest(
                    {
                        "id": document_id,
                        "source_type": ".pdf",
                        "sections": [toc["id"], chapter["id"]],
                        "pdf_bookmarks": [
                            {"title": "Introduction", "level": 0, "page": 2},
                            {"title": "User Guide", "level": 0, "page": 3},
                            {"title": "Developer Guide", "level": 0, "page": 4},
                        ],
                    }
                )

                skipped = document_workspace.table_of_contents_section_ids(
                    document_id, [toc["id"], chapter["id"]]
                )

        self.assertEqual(skipped, [toc["id"]])

    def test_detects_strong_epub_table_of_contents_sections(self):
        with tempfile.TemporaryDirectory() as directory:
            documents = Path(directory) / "documents"
            with patch.object(document_workspace, "DOCUMENTS_DIR", documents):
                document_id = "epub_toc"
                (documents / document_id).mkdir(parents=True)
                contents = document_workspace._new_section(
                    "Contents",
                    "Opening . . . 1\nGetting started . . . 5\nAdvanced use . . . 9",
                    chapter_title="Contents",
                    chapter_level=0,
                )
                chapters = [
                    document_workspace._new_section(
                        title, "Chapter text.", chapter_title=title, chapter_level=0
                    )
                    for title in ("Opening", "Getting started", "Advanced use")
                ]
                for section in [contents, *chapters]:
                    document_workspace.save_section(document_id, section)
                document_workspace.save_manifest(
                    {
                        "id": document_id,
                        "source_type": ".epub",
                        "sections": [contents["id"], *(section["id"] for section in chapters)],
                    }
                )

                skipped = document_workspace.table_of_contents_section_ids(
                    document_id, [contents["id"], *(section["id"] for section in chapters)]
                )

        self.assertEqual(skipped, [contents["id"]])

    def test_epub_navigation_anchors_split_one_spine_item(self):
        chapter = EpubChapter(
            "book.xhtml",
            b"<html><body><h1 id='first'>First</h1><p>One.</p><h1 id='second'>Second</h1><p>Two.</p></body></html>",
        )
        navigation = [
            EpubNavigation("First", 0, "book.xhtml", "first"),
            EpubNavigation("Second", 1, "book.xhtml", "second"),
        ]
        with patch.object(
            document_workspace,
            "read_epub_document",
            return_value=([chapter], navigation),
        ):
            sections = document_workspace._extract_epub(Path("book.epub"))

        self.assertEqual([section["title"] for section in sections], ["First", "Second"])
        self.assertEqual([section["chapter_level"] for section in sections], [0, 1])
        self.assertIn("One.", sections[0]["text"])
        self.assertNotIn("Two.", sections[0]["text"])

    def test_epub_navigation_anchors_preserve_utf16_content(self):
        content = "<html><body><h1 id='first'>First</h1><p>One.</p><h1 id='second'>Second</h1><p>Two.</p></body></html>".encode("utf-16")
        navigation = [
            EpubNavigation("First", 0, "book.xhtml", "first"),
            EpubNavigation("Second", 0, "book.xhtml", "second"),
        ]

        parts = document_workspace._split_epub_navigation(content, navigation)

        self.assertEqual(len(parts), 3)  # The markup before the first anchor is retained.
        self.assertIn("First", parts[1][1].decode("utf-8"))
        self.assertIn("Second", parts[2][1].decode("utf-8"))

    def test_unmatched_epub_anchor_keeps_preceding_content(self):
        chapter = EpubChapter(
            "book.xhtml",
            b"<html><body><p>Intro.</p><h1 id='second'>Second</h1><p>Two.</p></body></html>",
        )
        navigation = [
            EpubNavigation("Opening", 0, "book.xhtml", "missing"),
            EpubNavigation("Second", 0, "book.xhtml", "second"),
        ]
        with patch.object(
            document_workspace,
            "read_epub_document",
            return_value=([chapter], navigation),
        ):
            sections = document_workspace._extract_epub(Path("book.epub"))

        self.assertEqual([section["title"] for section in sections], ["Opening", "Second"])
        self.assertIn("Intro.", sections[0]["text"])

    def test_document_source_rejects_project_traversal(self):
        with tempfile.TemporaryDirectory() as directory:
            with patch.object(document_workspace, "DOCUMENTS_DIR", Path(directory) / "documents"):
                with self.assertRaises(VoxBenchError):
                    document_workspace.document_source_path("../../voxbench.json")

    def test_selects_an_inclusive_pdf_page_range(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            documents = root / "documents"
            with patch.object(document_workspace, "DOCUMENTS_DIR", documents):
                document_id = "page_range"
                (documents / document_id).mkdir(parents=True)
                sections = [
                    document_workspace._new_section(
                        f"Page {page}", f"Text {page}", source_page=page
                    )
                    for page in range(1, 6)
                ]
                for section in sections:
                    document_workspace.save_section(document_id, section)
                document_workspace.save_manifest(
                    {
                        "id": document_id,
                        "sections": [item["id"] for item in sections],
                    }
                )

                selected = document_workspace.section_ids_in_page_range(
                    document_id, 2, 4
                )

                self.assertEqual(selected, [item["id"] for item in sections[1:4]])

    def test_applies_cleanup_to_selected_sections(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            documents = root / "documents"
            with patch.object(document_workspace, "DOCUMENTS_DIR", documents):
                document_id = "batch_cleanup"
                (documents / document_id).mkdir(parents=True)
                sections = [
                    document_workspace._new_section("One", "First\nline"),
                    document_workspace._new_section("Two", "Second\nline"),
                ]
                for section in sections:
                    document_workspace.save_section(document_id, section)
                document_workspace.save_manifest(
                    {
                        "id": document_id,
                        "sections": [item["id"] for item in sections],
                    }
                )

                changed = document_workspace.apply_cleanup_to_sections(
                    document_id,
                    [sections[1]["id"]],
                    "Join broken lines",
                )

                self.assertEqual(changed, 1)
                self.assertEqual(
                    document_workspace.load_section(document_id, sections[0]["id"])[
                        "text"
                    ],
                    "First\nline",
                )
                self.assertEqual(
                    document_workspace.load_section(document_id, sections[1]["id"])[
                        "text"
                    ],
                    "Second line",
                )

    def test_cleans_safe_extraction_artifacts(self):
        cleaned = document_workspace.clean_text(
            "\u2022 Arch\u2019s\u00a0guide\ufffe\n\u2022 Next item",
            "Clean extraction artifacts",
        )
        self.assertEqual(cleaned, "Arch's guide\nNext item")

    def test_ignores_an_empty_section_until_it_contains_text(self):
        with tempfile.TemporaryDirectory() as directory:
            documents = Path(directory) / "documents"
            with patch.object(document_workspace, "DOCUMENTS_DIR", documents):
                document_id = "ignore_empty"
                (documents / document_id).mkdir(parents=True)
                section = document_workspace._new_section("Empty", "")
                document_workspace.save_section(document_id, section)
                document_workspace.save_manifest(
                    {"id": document_id, "sections": [section["id"]]}
                )

                self.assertEqual(
                    document_workspace.set_empty_sections_ignored(document_id, True), 1
                )
                document_workspace.save_editor_section(document_id, section["id"], "Empty", "")
                self.assertEqual(
                    document_workspace.load_section(document_id, section["id"])["status"], "Skipped"
                )

                self.assertEqual(
                    document_workspace.set_empty_sections_ignored(document_id, False), 1
                )
                self.assertEqual(
                    document_workspace.load_section(document_id, section["id"])["status"], "Needs review"
                )

                document_workspace.save_editor_section(document_id, section["id"], "Empty", "Has text")
                self.assertEqual(
                    document_workspace.load_section(document_id, section["id"])["status"], "Ready"
                )

    def test_removes_headers_only_from_selected_sections(self):
        with tempfile.TemporaryDirectory() as directory:
            documents = Path(directory) / "documents"
            with patch.object(document_workspace, "DOCUMENTS_DIR", documents):
                document_id = "selected_headers"
                (documents / document_id).mkdir(parents=True)
                sections = [
                    document_workspace._new_section("One", "Header\nFirst"),
                    document_workspace._new_section("Two", "Header\nSecond"),
                    document_workspace._new_section("Three", "Header\nThird"),
                ]
                for section in sections:
                    document_workspace.save_section(document_id, section)
                document_workspace.save_manifest(
                    {"id": document_id, "sections": [section["id"] for section in sections]}
                )

                changed = document_workspace.remove_repeated_headers_footers(
                    document_id, [sections[0]["id"], sections[1]["id"]]
                )

                self.assertEqual(changed, 2)
                self.assertEqual(
                    document_workspace.load_section(document_id, sections[2]["id"])["text"],
                    "Header\nThird",
                )
    def test_imports_docx_headings_and_tables_as_sections(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "sample.docx"
            document = Document()
            document.core_properties.title = "Sample document"
            document.add_heading("Opening", level=1)
            document.add_paragraph("Hello from Word.")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Speaker"
            table.cell(0, 1).text = "Line"
            document.add_heading("Closing", level=1)
            document.add_paragraph("The end.")
            document.save(source)

            with patch.object(
                document_workspace,
                "DOCUMENTS_DIR",
                root / "documents",
            ):
                document_id = document_workspace.import_document(str(source))
                manifest = document_workspace.load_manifest(document_id)
                self.assertEqual(manifest["source_type"], ".docx")
                self.assertEqual(len(manifest["sections"]), 2)
                first = document_workspace.load_section(
                    document_id,
                    manifest["sections"][0],
                )
                second = document_workspace.load_section(
                    document_id,
                    manifest["sections"][1],
                )
                self.assertEqual(first["title"], "Opening")
                self.assertIn("Hello from Word.", first["text"])
                self.assertIn("Speaker | Line", first["text"])
                self.assertIn("<table>", first["source_html"])
                self.assertEqual(second["title"], "Closing")
                self.assertIn("The end.", second["text"])

    def test_imports_pdf_as_page_sections(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "sample.pdf"
            write_blank_pdf(source, page_count=2)

            with patch.object(
                document_workspace,
                "DOCUMENTS_DIR",
                root / "documents",
            ):
                document_id = document_workspace.import_document(
                    str(source), source_name="Original Book.pdf"
                )
                manifest = document_workspace.load_manifest(document_id)
                self.assertEqual(manifest["source_type"], ".pdf")
                self.assertEqual(manifest["source_name"], "Original Book.pdf")
                self.assertEqual(len(manifest["sections"]), 2)
                first = document_workspace.load_section(
                    document_id,
                    manifest["sections"][0],
                )
                self.assertEqual(first["title"], "Page 1")
                self.assertEqual(first["source_page"], 1)
                self.assertEqual(document_workspace.document_page_count(document_id), 2)

    def test_imports_epub_spine_and_preserves_source_html(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "sample.epub"
            write_epub(source)

            with patch.object(
                document_workspace,
                "DOCUMENTS_DIR",
                root / "documents",
            ):
                document_id = document_workspace.import_document(str(source))
                manifest = document_workspace.load_manifest(document_id)
                self.assertEqual(len(manifest["sections"]), 2)
                first = document_workspace.load_section(
                    document_id,
                    manifest["sections"][0],
                )
                second = document_workspace.load_section(
                    document_id,
                    manifest["sections"][1],
                )
                self.assertEqual(first["title"], "Opening")
                self.assertIn("Hello document world.", first["text"])
                self.assertNotIn("<script", first["source_html"])
                self.assertEqual(second["title"], "Next chapter")
                self.assertIn("Second section.", second["text"])

    def test_edit_restore_search_and_restructure(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            documents = root / "documents"
            with patch.object(document_workspace, "DOCUMENTS_DIR", documents):
                document_id = "test_document"
                project = documents / document_id
                project.mkdir(parents=True)
                sections = [
                    document_workspace._new_section("One", "Hello world."),
                    document_workspace._new_section("Two", "Second section."),
                ]
                for section in sections:
                    document_workspace.save_section(document_id, section)
                document_workspace.save_manifest(
                    {
                        "id": document_id,
                        "title": "Test",
                        "source_name": "test.epub",
                        "source_path": str(project / "source.epub"),
                        "source_type": ".epub",
                        "created_at": "2026-01-01T00:00:00",
                        "sections": [item["id"] for item in sections],
                    }
                )

                document_workspace.save_editor_section(
                    document_id,
                    sections[0]["id"],
                    "Renamed",
                    "Hello edited world.",
                )
                count = document_workspace.replace_text(
                    document_id,
                    sections[0]["id"],
                    "world",
                    "book",
                    "Entire document",
                )
                self.assertEqual(count, 1)
                duplicate = document_workspace.restructure_section(
                    document_id,
                    sections[0]["id"],
                    "Duplicate",
                )
                self.assertEqual(
                    len(document_workspace.load_manifest(document_id)["sections"]),
                    3,
                )
                document_workspace.restructure_section(
                    document_id,
                    duplicate,
                    "Remove",
                )
                restored, status = document_workspace.restore_section(
                    document_id,
                    sections[0]["id"],
                )
                self.assertEqual(restored, "Hello world.")
                self.assertEqual(status, "Needs review")

    def test_entire_document_cleanup_skips_empty_sections(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            documents = root / "documents"
            with patch.object(document_workspace, "DOCUMENTS_DIR", documents):
                document_id = "cleanup_test"
                (documents / document_id).mkdir(parents=True)
                sections = [
                    document_workspace._new_section(
                        "One",
                        "Repeated header\nA hyphen-\nated line.\nRepeated footer",
                    ),
                    document_workspace._new_section(
                        "Two",
                        "Repeated header\n\nRepeated footer",
                    ),
                ]
                for section in sections:
                    document_workspace.save_section(document_id, section)
                document_workspace.save_manifest(
                    {
                        "id": document_id,
                        "title": "Cleanup",
                        "source_name": "cleanup.pdf",
                        "source_path": str(documents / document_id / "source.pdf"),
                        "source_type": ".pdf",
                        "created_at": "2026-01-01T00:00:00",
                        "sections": [item["id"] for item in sections],
                    }
                )
                ready = document_workspace.prepare_entire_document(document_id)
                self.assertEqual(ready, [sections[0]["id"]])
                first = document_workspace.load_section(
                    document_id,
                    sections[0]["id"],
                )
                second = document_workspace.load_section(
                    document_id,
                    sections[1]["id"],
                )
                self.assertEqual(first["text"], "A hyphenated line.")
                self.assertEqual(second["status"], "Skipped")

    def test_document_audio_is_saved_under_outputs(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            documents = root / "documents"
            outputs = root / "outputs"
            with (
                patch.object(document_workspace, "DOCUMENTS_DIR", documents),
                patch.object(document_workspace, "OUTPUTS_DIR", outputs),
            ):
                document_id = "20260101_120000_test"
                (documents / document_id).mkdir(parents=True)
                section = document_workspace._new_section("Opening", "Hello")
                document_workspace.save_section(document_id, section)
                document_workspace.save_manifest(
                    {
                        "id": document_id,
                        "title": "Test book",
                        "source_name": "test.epub",
                        "source_path": str(documents / document_id / "source.epub"),
                        "source_type": ".epub",
                        "created_at": "2026-01-01T12:00:00",
                        "sections": [section["id"]],
                    }
                )
                target = document_workspace.save_document_audio(
                    document_id,
                    section["id"],
                    np.zeros(2400, dtype=np.float32),
                    24000,
                )
                self.assertEqual(target.parent, outputs)
                self.assertTrue(target.is_file())
                document_workspace.clear_document_audio_paths(document_id, [section["id"]])
                self.assertIsNone(
                    document_workspace.load_section(document_id, section["id"])["audio_path"]
                )

    def test_existing_document_audio_is_moved_into_outputs(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            documents = root / "documents"
            outputs = root / "outputs"
            with (
                patch.object(document_workspace, "DOCUMENTS_DIR", documents),
                patch.object(document_workspace, "OUTPUTS_DIR", outputs),
            ):
                document_id = "20260101_120000_test"
                (documents / document_id).mkdir(parents=True)
                section = document_workspace._new_section("Opening", "Hello")
                document_workspace.save_section(document_id, section)
                document_workspace.save_manifest(
                    {
                        "id": document_id,
                        "title": "Test book",
                        "source_name": "test.epub",
                        "source_path": str(documents / document_id / "source.epub"),
                        "source_type": ".epub",
                        "created_at": "2026-01-01T12:00:00",
                        "sections": [section["id"]],
                    }
                )
                source = root / "section.wav"
                sf.write(source, np.zeros(2400, dtype=np.float32), 24000)
                target = document_workspace.save_document_audio(
                    document_id, section["id"], source, 24000
                )
                self.assertTrue(target.is_file())
                self.assertFalse(source.exists())

    def test_clearing_documents_preserves_generated_outputs(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            documents = root / "documents"
            project = documents / "stored_document"
            project.mkdir(parents=True)
            (project / "document.json").write_text("{}", encoding="utf-8")
            outputs = root / "outputs"
            outputs.mkdir()
            generated = outputs / "chapter.wav"
            generated.write_bytes(b"generated audio")

            with (
                patch.object(document_workspace, "PROJECT_DIR", root),
                patch.object(document_workspace, "DOCUMENTS_DIR", documents),
            ):
                removed = document_workspace.clear_document_projects()

            self.assertEqual(removed, 1)
            self.assertFalse(documents.exists())
            self.assertTrue(generated.is_file())


if __name__ == "__main__":
    unittest.main()
